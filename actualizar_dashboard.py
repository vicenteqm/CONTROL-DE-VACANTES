#!/usr/bin/env python3
"""Regenera data.js con los Excel actuales y, opcionalmente, hace commit/push a GitHub."""
from __future__ import annotations
import fnmatch
import json
import logging
import os
import subprocess
import sys
from datetime import datetime
from pathlib import Path
import unicodedata

try:
    from openpyxl import load_workbook
except ImportError:
    print("Falta la dependencia openpyxl. Ejecuta instalar.bat y vuelve a intentar.")
    raise SystemExit(2)

ROOT = Path(__file__).resolve().parent
CONFIG_PATH = ROOT / "config.json"
LOG_DIR = ROOT / "logs"
LOG_DIR.mkdir(exist_ok=True)
LOG_PATH = LOG_DIR / "actualizacion.log"
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)s | %(message)s",
    handlers=[logging.FileHandler(LOG_PATH, encoding="utf-8"), logging.StreamHandler(sys.stdout)],
)

VALID_PROFILES = {"PROMOTOR DE CREDITO", "SUBLIDER"}
CODE_LABELS = {
    "A": "Asistencia", "F": "Falta", "R": "Retardo", "I": "Incapacidad",
    "V": "Vacaciones", "D": "Descanso", "OTROS": "Otros códigos",
}
CODE_ORDER = ["A", "F", "R", "I", "V", "D", "OTROS"]


def text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value).replace("\u00a0", " ")
    return " ".join(s.strip().upper().split())


def normalized(value) -> str:
    s = text(value)
    return "".join(ch for ch in unicodedata.normalize("NFD", s) if unicodedata.category(ch) != "Mn")


def det_key(value) -> str:
    return text(value)


def store_key(socio, determinante) -> str:
    socio_norm = normalized(socio)
    det = det_key(determinante)
    if not socio_norm or not det:
        return ""
    return f"{socio_norm} {det}"


def number(value, default=0) -> int:
    if value is None or value == "":
        return default
    try:
        n = float(value)
        if n < 0 or n > 100000:
            return default
        return int(n)
    except (TypeError, ValueError):
        return default


def find_col(headers, wanted, occurrence=0):
    wanted = normalized(wanted)
    matching = [i for i, h in enumerate(headers) if normalized(h) == wanted]
    if len(matching) <= occurrence:
        raise ValueError(f"No se encontró la columna obligatoria '{wanted}'. Encabezados detectados: {headers[:30]}")
    return matching[occurrence]


def find_header_row(ws, required_headers, scan_rows=8):
    required = {normalized(x) for x in required_headers}
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(scan_rows, ws.max_row), values_only=True), start=1):
        headers = [normalized(v) for v in row]
        if required.issubset(set(headers)):
            return row_idx, list(row)
    raise ValueError(f"No se localizaron los encabezados esperados en la hoja '{ws.title}'.")


def latest_file(folder_value, pattern, description):
    folder = Path(os.path.expandvars(os.path.expanduser(folder_value))).resolve()
    if not folder.is_dir():
        raise FileNotFoundError(f"No existe la carpeta de {description}: {folder}\nRevisa config.json o vuelve a ejecutar configurar_carpetas.bat.")
    candidates = [p for p in folder.iterdir() if p.is_file() and p.suffix.lower() == ".xlsx"
                  and not p.name.startswith("~$") and fnmatch.fnmatch(p.name.lower(), pattern.lower())]
    if not candidates:
        raise FileNotFoundError(f"No se encontró un Excel que coincida con '{pattern}' en {folder}.")
    return max(candidates, key=lambda p: p.stat().st_mtime)


def read_base(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    expected = {"SOCIO", "NUMERO DE TIENDA/DETERMINANTE", "TIENDA", "ESTATUS", "PROMOTORES REQUERIDOS"}
    selected = None
    for ws in wb.worksheets:
        try:
            header_row, headers = find_header_row(ws, expected)
            selected = (ws, header_row, headers)
            break
        except ValueError:
            continue
    if selected is None:
        wb.close()
        raise ValueError("No encontré una hoja de base de tiendas con los encabezados requeridos.")
    ws, header_row, headers = selected
    # La columna SUBLIDER requerida es la primera aparición de ese encabezado.
    cols = {
        "socio": find_col(headers, "SOCIO"),
        "det": find_col(headers, "NUMERO DE TIENDA/DETERMINANTE"),
        "tienda": find_col(headers, "TIENDA"),
        "estatus": find_col(headers, "ESTATUS"),
        "prom_req": find_col(headers, "PROMOTORES REQUERIDOS"),
        "sub_req": find_col(headers, "SUBLIDER", 0),
        "ciudad": find_col(headers, "CIUDAD"),
        "estado": find_col(headers, "ESTADO"),
        "region": find_col(headers, "REGION"),
        "zona": find_col(headers, "ZONAS"),
    }
    stores = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        socio, det = row[cols["socio"]], row[cols["det"]]
        key = store_key(socio, det)
        if not key:
            continue
        prom_req = number(row[cols["prom_req"]])
        sub_req = number(row[cols["sub_req"]])
        stores.append({
            "key": key,
            "socio": text(socio) or "SIN DATO",
            "determinante": det_key(det),
            "tienda": text(row[cols["tienda"]]) or "SIN NOMBRE",
            "ciudad": text(row[cols["ciudad"]]),
            "estado": text(row[cols["estado"]]),
            "region": text(row[cols["region"]]) or "SIN REGIÓN",
            "zona": text(row[cols["zona"]]),
            "estatus": normalized(row[cols["estatus"]]) or "SIN ESTATUS",
            "promReq": prom_req,
            "subReq": sub_req,
            "totalReq": prom_req + sub_req,
        })
    wb.close()
    if not stores:
        raise ValueError("La base de tiendas no contiene filas de datos utilizables.")
    return stores


def attendance_code(value):
    code = normalized(value)
    if not code or code in {"0", "N/A", "NA", "-", "--"}:
        return ""
    if code in CODE_LABELS and code != "OTROS":
        return code
    # Los códigos no reconocidos se agrupan en OTROS, sin inventar su significado.
    if len(code) <= 4:
        return "OTROS"
    return ""


def read_attendance(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    expected = {"SOCIO", "DETERMINANTE", "ESTATUS", "PERFIL DE PUESTO BANCO"}
    selected = None
    for ws in wb.worksheets:
        try:
            header_row, headers = find_header_row(ws, expected)
            selected = (ws, header_row, headers)
            break
        except ValueError:
            continue
    if selected is None:
        wb.close()
        raise ValueError("No encontré una hoja de asistencia con los encabezados requeridos.")
    ws, header_row, headers = selected
    cols = {
        "socio": find_col(headers, "SOCIO"),
        "det": find_col(headers, "DETERMINANTE"),
        "estatus": find_col(headers, "ESTATUS"),
        "profile": find_col(headers, "PERFIL DE PUESTO BANCO"),
    }
    day_columns = {}
    for idx, value in enumerate(headers):
        # En el archivo de origen los días son encabezados numéricos (1, 2, 3...).
        day = None
        if isinstance(value, (int, float)) and float(value).is_integer():
            day = int(value)
        elif isinstance(value, str) and value.strip().isdigit():
            day = int(value.strip())
        if day is not None and 1 <= day <= 31:
            day_columns[day] = idx
    if not day_columns:
        wb.close()
        raise ValueError("No encontré columnas de días (1–31) en la hoja de asistencia.")

    rows = []
    all_keys = set()
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        key = store_key(row[cols["socio"]], row[cols["det"]])
        if not key:
            continue
        all_keys.add(key)
        status = normalized(row[cols["estatus"]])
        profile = normalized(row[cols["profile"]])
        if status != "ACTIVO" or profile not in VALID_PROFILES:
            continue
        codes = {}
        for day, col_idx in day_columns.items():
            raw = row[col_idx] if col_idx < len(row) else None
            code = attendance_code(raw)
            if code:
                codes[day] = code
        rows.append({"key": key, "profile": profile, "codes": codes})
    wb.close()
    return rows, all_keys


def build_data(base_file, attendance_file):
    stores = read_base(base_file)
    eligible_all, all_attendance_keys = read_attendance(attendance_file)
    operating_keys = {store["key"] for store in stores if store["estatus"] == "OPERANDO"}
    # Cobertura y gráfica diaria: solo empleados elegibles ligados a tiendas OPERANDO.
    eligible_rows = [emp for emp in eligible_all if emp["key"] in operating_keys]
    by_store = {}
    by_day = {}
    totals = {code: 0 for code in CODE_ORDER}
    for emp in eligible_rows:
        bucket = by_store.setdefault(emp["key"], {"PROMOTOR DE CREDITO": 0, "SUBLIDER": 0})
        bucket[emp["profile"]] += 1
        for day, code in emp["codes"].items():
            by_day.setdefault(day, {c: 0 for c in CODE_ORDER})[code] += 1
            totals[code] += 1

    for store in stores:
        count = by_store.get(store["key"], {})
        store["promActive"] = count.get("PROMOTOR DE CREDITO", 0)
        store["subActive"] = count.get("SUBLIDER", 0)
        store["totalActive"] = store["promActive"] + store["subActive"]
        store["promGap"] = max(0, store["promReq"] - store["promActive"])
        store["subGap"] = max(0, store["subReq"] - store["subActive"])
        store["gap"] = store["promGap"] + store["subGap"]
        store["matched"] = store["key"] in all_attendance_keys
        store["hasEligible"] = store["key"] in by_store
        del store["key"]

    operating = [s for s in stores if s["estatus"] == "OPERANDO"]
    non_operating = [s for s in stores if s["estatus"] != "OPERANDO"]
    required = sum(s["totalReq"] for s in operating)
    active = sum(s["totalActive"] for s in operating)
    role_gap = sum(s["gap"] for s in operating)
    status_summary = {}
    for s in stores:
        status_summary[s["estatus"]] = status_summary.get(s["estatus"], 0) + 1

    regions = {}
    for s in operating:
        region = s["region"] or "SIN REGIÓN"
        item = regions.setdefault(region, {"region": region, "required": 0, "active": 0, "gap": 0, "stores": 0})
        item["required"] += s["totalReq"]
        item["active"] += s["totalActive"]
        item["gap"] += s["gap"]
        item["stores"] += 1

    # Oculta días futuros/vacíos: un día solo entra al gráfico si hay algún código capturado.
    attendance_by_day = []
    for day in sorted(by_day):
        counts = by_day[day]
        if sum(counts.values()) > 0:
            attendance_by_day.append({"day": day, "counts": counts})
    present_codes = [c for c in CODE_ORDER if totals[c] > 0]
    return {
        "meta": {
            "source": f"{base_file.name} + {attendance_file.name}",
            "scope": "Tiendas OPERANDO; asistencia con ESTATUS ACTIVO y perfil PROMOTOR DE CREDITO o SUBLIDER",
            "generated": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "generatedDisplay": datetime.now().strftime("%d/%m/%Y %H:%M"),
            "attendanceDays": [x["day"] for x in attendance_by_day],
            "attendanceCodes": {k: CODE_LABELS[k] for k in present_codes},
        },
        "stores": stores,
        "summary": {
            "storesOperating": len(operating),
            "required": required,
            "active": active,
            "gap": role_gap,
            "netGap": required - active,
            "coverage": round(active / required * 100, 1) if required else 0,
            "promRequired": sum(s["promReq"] for s in operating),
            "subRequired": sum(s["subReq"] for s in operating),
            "promActive": sum(s["promActive"] for s in operating),
            "subActive": sum(s["subActive"] for s in operating),
            "nonOperating": len(non_operating),
            "matchedStores": sum(s["matched"] for s in operating),
            "unmatchedStores": sum(not s["matched"] for s in operating),
            "eligibleEmployees": len(eligible_rows),
            "operatingStoresWithoutEligible": sum(not s["hasEligible"] for s in operating),
        },
        "statusSummary": status_summary,
        "regions": list(regions.values()),
        "attendanceByDay": attendance_by_day,
        "attendanceTotals": totals,
        "attendanceCodes": {k: CODE_LABELS[k] for k in CODE_ORDER},
        "unmatched": [
            {"socio": s["socio"], "determinante": s["determinante"], "tienda": s["tienda"], "ciudad": s["ciudad"], "region": s["region"]}
            for s in operating if not s["matched"]
        ],
    }


def git_publish():
    git_dir = ROOT / ".git"
    if not git_dir.exists():
        logging.warning("La carpeta no es un clon Git (no existe .git). Se generó data.js, pero no se publicó. Coloca el proyecto dentro del repositorio clonado.")
        return False
    def run_git(args, check=True):
        proc = subprocess.run(["git", *args], cwd=ROOT, text=True, capture_output=True, encoding="utf-8", errors="replace")
        if proc.stdout.strip(): logging.info(proc.stdout.strip())
        if proc.stderr.strip(): logging.info(proc.stderr.strip())
        if check and proc.returncode != 0:
            raise RuntimeError(f"Falló el comando git {' '.join(args)}: {proc.stderr.strip() or proc.stdout.strip()}")
        return proc
    # Solo agrega data.js; no incluye los Excel, config local, ni otros cambios del usuario.
    run_git(["add", "--", "data.js"])
    diff = run_git(["diff", "--cached", "--quiet", "--", "data.js"], check=False)
    if diff.returncode == 0:
        logging.info("No hay cambios nuevos en data.js.")
        # Si hubo un push fallido anterior, reintenta enviar los commits locales pendientes.
        ahead = run_git(["rev-list", "--count", "@{u}..HEAD"], check=False)
        if ahead.returncode == 0 and ahead.stdout.strip().isdigit() and int(ahead.stdout.strip()) > 0:
            run_git(["push"])
        return True
    if diff.returncode != 1:
        raise RuntimeError("No fue posible comprobar los cambios preparados por Git.")
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    run_git(["commit", "--only", "data.js", "-m", f"Actualización diaria de datos INVEX - {stamp}"])
    run_git(["push"])
    logging.info("GitHub recibió la actualización correctamente.")
    return True


def main():
    if not CONFIG_PATH.exists():
        raise FileNotFoundError("No existe config.json. Ejecuta configurar_carpetas.bat para seleccionar tus carpetas de Excel.")
    config = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
    base_file = latest_file(config["base_folder"], config.get("base_pattern", "*BASE*TIENDAS*.xlsx"), "base de tiendas")
    attendance_file = latest_file(config["attendance_folder"], config.get("attendance_pattern", "*ASISTENCIA*.xlsx"), "asistencia")
    logging.info("Base seleccionada: %s", base_file)
    logging.info("Asistencia seleccionada: %s", attendance_file)
    if base_file.resolve() == attendance_file.resolve():
        raise ValueError("Las carpetas seleccionadas apuntan al mismo archivo. Revisa la configuración.")
    data = build_data(base_file, attendance_file)
    output = ROOT / "data.js"
    # Si no cambiaron datos ni archivos fuente, evita commits duplicados por ejecuciones repetidas.
    if output.exists():
        try:
            previous_text = output.read_text(encoding="utf-8")
            prefix = "window.INVEX_DATA = "
            previous_json = previous_text[len(prefix):].strip().removesuffix(";") if previous_text.startswith(prefix) else ""
            previous = json.loads(previous_json) if previous_json else None
            def comparable(obj):
                obj = json.loads(json.dumps(obj, ensure_ascii=False))
                if isinstance(obj.get("meta"), dict):
                    obj["meta"].pop("generated", None)
                    obj["meta"].pop("generatedDisplay", None)
                return obj
            if previous is not None and comparable(previous) == comparable(data):
                logging.info("No cambiaron los indicadores ni la información agregada. No se genera un commit nuevo.")
                if config.get("auto_push", True):
                    git_publish()  # Reintenta un push pendiente si una ejecución previa falló.
                return
        except (ValueError, json.JSONDecodeError):
            logging.warning("No fue posible comparar el data.js anterior; se regenerará.")
    serialized = "window.INVEX_DATA = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";\n"
    tmp = ROOT / "data.js.tmp"
    tmp.write_text(serialized, encoding="utf-8")
    tmp.replace(output)
    s = data["summary"]
    logging.info("Datos generados: %s tiendas operando | plantilla %s | activos %s | brecha por perfil %s | cobertura %s%%", s["storesOperating"], s["required"], s["active"], s["gap"], s["coverage"])
    logging.info("Días de asistencia con códigos capturados: %s", ", ".join(map(str, data["meta"]["attendanceDays"])) or "ninguno")
    if config.get("auto_push", True):
        git_publish()
    else:
        logging.info("auto_push está desactivado; data.js se actualizó localmente sin publicar.")

if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        logging.exception("ACTUALIZACIÓN FALLIDA: %s", exc)
        print("\nNo se completó la actualización. Consulta logs/actualizacion.log para revisar el detalle.")
        if "--scheduled" not in sys.argv and sys.stdin is not None and sys.stdin.isatty():
            input("Presiona Enter para cerrar...")
        raise SystemExit(1)
